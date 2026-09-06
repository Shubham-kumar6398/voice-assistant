"""
api/routes.py — REST endpoints: system stats, recordings, to-dos and RAG documents.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Optional

import psutil
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response

from config import RECORDINGS_DIR, TODO_STATUSES

router = APIRouter()

_SESSION_RE = re.compile(r"^\d{8}_\d{6}_\d{1,10}$")
_DOC_ID_RE = re.compile(r"^[0-9a-f]{8}$")


# ── System stats ──────────────────────────────────────────────────────────────

@router.get("/stats")
async def get_stats():
    """Live CPU / RAM / disk usage."""
    vm = psutil.virtual_memory()
    return {
        "cpu_pct":      round(psutil.cpu_percent(interval=None), 1),
        "ram_used_gb":  round(vm.used / (1024 ** 3), 2),
        "ram_total_gb": round(vm.total / (1024 ** 3), 2),
        "ram_pct":      round(vm.percent, 1),
        "disk_pct":     round(psutil.disk_usage("/").percent, 1),
    }


# ── Recordings ────────────────────────────────────────────────────────────────

@router.get("/recordings/{filename}")
async def get_recording(filename: str):
    if "/" in filename or ".." in filename or not filename.endswith(".wav"):
        return Response(status_code=400)
    path = Path(RECORDINGS_DIR) / filename
    if not path.exists():
        return Response(status_code=404)
    return FileResponse(str(path), media_type="audio/wav")


@router.delete("/recordings/session/{session_prefix}")
async def delete_session_recordings(session_prefix: str):
    """Delete every WAV for one session id (YYYYMMDD_HHMMSS_N)."""
    if not _SESSION_RE.match(session_prefix):
        return Response(status_code=400)
    deleted = []
    for f in Path(RECORDINGS_DIR).glob(f"{session_prefix}_*.wav"):
        f.unlink(missing_ok=True)
        deleted.append(f.name)
    return {"deleted": deleted}


# ── To-dos ────────────────────────────────────────────────────────────────────

def register_todo_routes(app, todo_store) -> None:
    """Attach the to-do routes, backed by the shared TodoStore."""

    @app.get("/todos")
    async def todos_list(status: Optional[str] = None):
        return todo_store.list(status)

    @app.post("/todos")
    async def todos_create(payload: dict):
        text = (payload.get("text") or "").strip()
        if not text:
            raise HTTPException(400, "text is required")
        return todo_store.create(text)

    @app.patch("/todos/{todo_id}")
    async def todos_update(todo_id: str, payload: dict):
        """Change a task's status, its wording, or both."""
        status = (payload.get("status") or "").strip()
        text = (payload.get("text") or "").strip()
        if not status and not text:
            raise HTTPException(400, "provide status and/or text")
        if status and status not in TODO_STATUSES:
            raise HTTPException(400, f"status must be one of {TODO_STATUSES}")

        todo = None
        if text:
            todo = todo_store.update_text(todo_id, text)
        if status:
            todo = todo_store.update_status(todo_id, status)
        if not todo:
            raise HTTPException(404, "Todo not found")
        return todo

    @app.delete("/todos/{todo_id}")
    async def todos_delete(todo_id: str):
        if not todo_store.delete(todo_id):
            raise HTTPException(404, "Todo not found")
        return {"deleted": todo_id}


# ── RAG documents ─────────────────────────────────────────────────────────────

def _extract_text(filename: str, content: bytes) -> str:
    """Plain text from a PDF, .docx or text upload. Raises ValueError if unsupported."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:      # tables hold real content in most SOPs
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    if ext == "doc":
        raise ValueError("Legacy .doc is not supported. Please save as .docx or PDF.")

    return content.decode("utf-8", errors="replace")


def register_rag_routes(app, rag_handler) -> None:
    """Attach the document routes, backed by the shared RAGHandler."""

    @app.post("/rag/upload")
    async def rag_upload(file: UploadFile = File(...)):
        if not rag_handler:
            return JSONResponse({"error": "RAG not enabled"}, status_code=503)
        filename = file.filename or "document"
        try:
            text = _extract_text(filename, await file.read())
        except ValueError as exc:
            return JSONResponse({"error": str(exc)}, status_code=400)
        except Exception as exc:
            return JSONResponse({"error": f"Could not read file: {exc}"}, status_code=400)
        if not text.strip():
            return JSONResponse({"error": "No text extracted"}, status_code=400)
        return rag_handler.add_document(text, filename)

    @app.get("/rag/documents")
    async def rag_list():
        return rag_handler.list_documents() if rag_handler else []

    @app.delete("/rag/documents/{doc_id}")
    async def rag_delete(doc_id: str):
        if not rag_handler:
            return JSONResponse({"error": "RAG not enabled"}, status_code=503)
        if not _DOC_ID_RE.match(doc_id):
            return Response(status_code=400)
        return {"deleted_chunks": rag_handler.delete_document(doc_id)}
